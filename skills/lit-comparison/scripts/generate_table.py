"""
Literature Comparison Table Generator

Produces a professionally formatted .xlsx file with three worksheets:
1. 文獻統整比較表 — Five-dimension comparison matrix
2. 交叉缺口矩陣 — Cross-gap matrix with color coding
3. 對研究的啟示 — Implications for the user's research

Usage:
    create_literature_comparison(
        papers=[...],           # list of paper metadata dicts
        dimensions=[...],       # list of 5 dimension content lists
        gaps=[...],             # list of gap row dicts
        implications=[...],     # list of implication row dicts
        output_path="output.xlsx",
        implication_sheet_title="對碩士論文的啟示"  # optional custom title
    )
"""

import argparse
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


# ── Style constants ──────────────────────────────────────────────────────────

HEADER_FILL = PatternFill('solid', fgColor='1F4E79')
ROW_HEADER_FILL = PatternFill('solid', fgColor='D6E4F0')
ALT_FILL = PatternFill('solid', fgColor='F2F7FB')
WHITE_FILL = PatternFill('solid', fgColor='FFFFFF')

GREEN_FILL = PatternFill('solid', fgColor='E2EFDA')
RED_FILL = PatternFill('solid', fgColor='FCE4EC')
YELLOW_FILL = PatternFill('solid', fgColor='FFF8E1')
GRAY_FILL = PatternFill('solid', fgColor='F5F5F5')
STAR_FILL = PatternFill('solid', fgColor='E8F5E9')
STAR2_FILL = PatternFill('solid', fgColor='C8E6C9')
MILD_YELLOW_FILL = PatternFill('solid', fgColor='FFF9C4')

HEADER_FONT = Font(name='Arial', bold=True, color='FFFFFF', size=11)
ROW_HEADER_FONT = Font(name='Arial', bold=True, color='1F4E79', size=11)
BODY_FONT = Font(name='Arial', size=10, color='333333')
BOLD_FONT = Font(name='Arial', size=10, bold=True)
STAR_FONT = Font(name='Arial', size=10, bold=True, color='2E7D32')
STAR2_FONT = Font(name='Arial', size=10, bold=True, color='1B5E20')
IMP_LABEL_FONT = Font(name='Arial', size=11, bold=True, color='1F4E79')

THIN_BORDER = Border(
    left=Side(style='thin', color='B0C4DE'),
    right=Side(style='thin', color='B0C4DE'),
    top=Side(style='thin', color='B0C4DE'),
    bottom=Side(style='thin', color='B0C4DE')
)

DIMENSION_LABELS = [
    "研究發展趨勢\n（研究焦點如何隨時間改變？）",
    "主要研究發現\n（核心主張與證據為何？）",
    "研究分歧與爭議\n（文獻之間的共識與衝突）",
    "文獻缺口\n（尚未被解決的問題）",
    "未來研究方向\n（根據缺口應研究什麼）"
]


def _col_width(n_papers):
    """Calculate column width based on number of papers."""
    if n_papers <= 3:
        return 52
    elif n_papers <= 5:
        return 42
    else:
        return 36


def _apply_gap_fill(cell, value, is_last_col=False):
    """Apply color fill to gap matrix cells based on markers."""
    val = str(value).lstrip()
    if is_last_col:
        if val.startswith('★★'):
            cell.fill = STAR2_FILL
            cell.font = STAR2_FONT
        elif val.startswith('★'):
            cell.fill = STAR_FILL
            cell.font = STAR_FONT
        elif val.startswith('✓'):
            cell.fill = MILD_YELLOW_FILL
        elif val.startswith('△'):
            cell.fill = YELLOW_FILL
        else:
            cell.fill = GRAY_FILL
    else:
        if val.startswith('✓'):
            cell.fill = GREEN_FILL
        elif val.startswith('✗'):
            cell.fill = RED_FILL
        elif val.startswith('△'):
            cell.fill = YELLOW_FILL
        elif val.startswith('—'):
            cell.fill = GRAY_FILL


def create_literature_comparison(
    papers,
    dimensions,
    gaps,
    implications,
    output_path,
    implication_sheet_title="對碩士論文的啟示"
):
    """
    Create a literature comparison Excel workbook.

    Args:
        papers: List of dicts with keys:
            - "header": Full header string for column (author, title, journal, details)
            - "short_name": Short label for gap matrix columns (e.g., "Tarsuslu\n(2024)")

        dimensions: List of 5 lists, each containing one string per paper.
            dimensions[0] = ["paper1 trend text", "paper2 trend text", ...]
            dimensions[1] = ["paper1 findings text", "paper2 findings text", ...]
            ... (research trends, key findings, disagreements, gaps, future directions)

        gaps: List of dicts with keys:
            - "label": Gap description (row header)
            - "cells": List of strings, one per paper + one for "your research" column
              Use markers: ✓, ✗, △, —, ★, ★★

        implications: List of dicts with keys:
            - "label": Implication category (e.g., "研究問題定位")
            - "content": Detailed content string

        output_path: Path to save the .xlsx file

        implication_sheet_title: Title for Sheet 3 (default: "對碩士論文的啟示")
    """
    # Validate before producing any file: zip() must never silently drop data.
    def text(value):
        if not isinstance(value, str) or len(value) > 32767 or ILLEGAL_CHARACTERS_RE.search(value):
            raise ValueError("Cells must be text, at most 32767 characters, without illegal control characters")
    if not isinstance(papers, list) or not 2 <= len(papers) <= 100:
        raise ValueError("Provide 2–100 papers; split larger comparisons into related groups")
    n_papers = len(papers)
    for paper in papers:
        text(paper["header"])
        text(paper["short_name"])
    if not isinstance(dimensions, list) or len(dimensions) != 5:
        raise ValueError("Exactly five dimensions are required")
    for row in dimensions:
        if not isinstance(row, list) or len(row) != n_papers:
            raise ValueError("Each dimension must contain one cell per paper")
        for value in row:
            text(value)
    if not isinstance(gaps, list) or not isinstance(implications, list):
        raise ValueError("Gaps and implications must be lists")
    for gap in gaps:
        text(gap["label"])
        if not isinstance(gap["cells"], list) or len(gap["cells"]) != n_papers + 1:
            raise ValueError("Each gap requires one cell per paper plus a research opportunity cell")
        for value in gap["cells"]:
            text(value)
    for implication in implications:
        text(implication["label"])
        text(implication["content"])
    if (not isinstance(implication_sheet_title, str) or not implication_sheet_title
            or len(implication_sheet_title) > 31
            or any(c in implication_sheet_title for c in '[]:*?/\\')
            or implication_sheet_title.casefold() in {"文獻統整比較表", "交叉缺口矩陣"}):
        raise ValueError("Invalid or duplicate implication sheet title")
    col_w = _col_width(n_papers)
    wb = Workbook()

    # ── Sheet 1: 文獻統整比較表 ──────────────────────────────────────────────

    ws = wb.active
    ws.title = "文獻統整比較表"

    # Headers
    headers = ["文獻統整面向"] + [p["header"] for p in papers]
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = THIN_BORDER

    # Body rows
    for row_idx, (label, dim_data) in enumerate(zip(DIMENSION_LABELS, dimensions), 2):
        row_values = [label] + list(dim_data)
        for col_idx, value in enumerate(row_values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            cell.border = THIN_BORDER
            if col_idx == 1:
                cell.font = ROW_HEADER_FONT
                cell.fill = ROW_HEADER_FILL
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            else:
                cell.font = BODY_FONT
                cell.fill = ALT_FILL if row_idx % 2 == 0 else WHITE_FILL

    # Column widths and row heights
    ws.column_dimensions['A'].width = 22
    for i in range(n_papers):
        col_letter = get_column_letter(i + 2)
        ws.column_dimensions[col_letter].width = col_w
    ws.row_dimensions[1].height = 75
    for r in range(2, 7):
        ws.row_dimensions[r].height = 300
    ws.freeze_panes = 'B2'

    # ── Sheet 2: 交叉缺口矩陣 ───────────────────────────────────────────────

    ws2 = wb.create_sheet("交叉缺口矩陣")

    gap_headers = ["研究缺口"] + [p["short_name"] for p in papers] + ["你的研究可切入？"]
    n_gap_cols = len(gap_headers)

    for col_idx, h in enumerate(gap_headers, 1):
        cell = ws2.cell(row=1, column=col_idx, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = THIN_BORDER

    for row_idx, gap in enumerate(gaps, 2):
        values = [gap["label"]] + gap["cells"]
        for col_idx, value in enumerate(values, 1):
            cell = ws2.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(
                horizontal='center' if col_idx > 1 else 'left',
                vertical='center', wrap_text=True
            )
            cell.border = THIN_BORDER
            cell.font = Font(name='Arial', size=10)

            if col_idx == 1:
                cell.font = BOLD_FONT
                cell.fill = ROW_HEADER_FILL
            elif col_idx == n_gap_cols:
                _apply_gap_fill(cell, value, is_last_col=True)
            else:
                _apply_gap_fill(cell, value, is_last_col=False)

    # Column widths
    ws2.column_dimensions['A'].width = 35
    for i in range(1, n_papers + 1):
        col_letter = get_column_letter(i + 1)
        ws2.column_dimensions[col_letter].width = 22 if n_papers <= 4 else 18
    last_col = get_column_letter(n_papers + 2)
    ws2.column_dimensions[last_col].width = 42
    ws2.row_dimensions[1].height = 40
    for r in range(2, 2 + len(gaps)):
        ws2.row_dimensions[r].height = 55
    ws2.freeze_panes = 'B2'

    # ── Sheet 3: 對研究的啟示 ────────────────────────────────────────────────

    ws3 = wb.create_sheet(implication_sheet_title)

    imp_headers = ["啟示面向", "具體內容"]
    for col_idx, h in enumerate(imp_headers, 1):
        cell = ws3.cell(row=1, column=col_idx, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = THIN_BORDER

    for row_idx, imp in enumerate(implications, 2):
        # Label column
        cell_label = ws3.cell(row=row_idx, column=1, value=imp["label"])
        cell_label.font = IMP_LABEL_FONT
        cell_label.fill = ROW_HEADER_FILL
        cell_label.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell_label.border = THIN_BORDER

        # Content column
        cell_content = ws3.cell(row=row_idx, column=2, value=imp["content"])
        cell_content.font = BODY_FONT
        cell_content.alignment = Alignment(vertical='center', wrap_text=True)
        cell_content.border = THIN_BORDER

    ws3.column_dimensions['A'].width = 20
    ws3.column_dimensions['B'].width = 92
    ws3.row_dimensions[1].height = 35
    for r in range(2, 2 + len(implications)):
        ws3.row_dimensions[r].height = 90

    # ── Save ─────────────────────────────────────────────────────────────────

    # Research content is literal text, never an executable Excel formula.
    for sheet in wb:
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
    output = Path(output_path)
    if output.suffix.lower() != ".xlsx":
        raise ValueError("Output must be an .xlsx file")
    output.parent.mkdir(parents=True, exist_ok=True)
    # Exclusive creation protects previous results and manual workbook edits.
    stream = output.open("xb")
    try:
        with stream:
            wb.save(stream)
    except Exception:
        output.unlink(missing_ok=True)
        raise
    print(f"Saved to {output_path}")
    return output_path


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Generate a literature comparison workbook from evidence-backed JSON")
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    data = json.loads(args.input.read_text(encoding="utf-8"))
    create_literature_comparison(
        papers=data["papers"], dimensions=data["dimensions"], gaps=data["gaps"],
        implications=data["implications"], output_path=args.output,
        implication_sheet_title=data.get("implication_sheet_title", "對研究的啟示"),
    )
