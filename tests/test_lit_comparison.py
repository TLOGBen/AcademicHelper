"""Synthetic workbook tests; no claims about actual research findings."""
import importlib.util
import json
from pathlib import Path
import subprocess
import sys

from openpyxl import load_workbook
import pytest

SCRIPT = Path(__file__).resolve().parents[1] / 'skills/lit-comparison/scripts/generate_table.py'
spec = importlib.util.spec_from_file_location('lit_generator', SCRIPT)
generator = importlib.util.module_from_spec(spec)
spec.loader.exec_module(generator)


def sample(n=2):
    return dict(
        papers=[dict(header=f'SYNTHETIC paper {i}', short_name=f'P{i}') for i in range(n)],
        dimensions=[[f'SYNTHETIC dimension {d}, P{i}; source pending' for i in range(n)] for d in range(5)],
        gaps=[dict(label='Synthetic candidate', cells=['？ not reported'] * n + ['△ requires evidence'])],
        implications=[dict(label='Synthetic advice', content='Not a real research finding')],
    )


def test_three_sheets_and_wide_columns(tmp_path):
    path = tmp_path / 'wide.xlsx'
    generator.create_literature_comparison(**sample(27), output_path=path)
    wb = load_workbook(path)
    assert wb.sheetnames == ['文獻統整比較表', '交叉缺口矩陣', '對碩士論文的啟示']
    ws = wb.worksheets[0]
    assert ws.max_column == 28 and ws.max_row == 6
    assert ws.column_dimensions['AB'].width == 36
    assert ws['AB1'].value == 'SYNTHETIC paper 26'
    assert ws.freeze_panes == 'B2'
    assert wb.worksheets[1].column_dimensions['AC'].width == 42
    assert ws['B2'].alignment.wrap_text
    wb.close()


def test_formula_looking_content_roundtrips_as_text(tmp_path):
    data = sample()
    text = '=HYPERLINK("https://example.invalid", "synthetic")'
    data['papers'][0]['header'] = text
    data['dimensions'][0][0] = text
    data['gaps'][0]['cells'][0] = text
    data['implications'][0]['content'] = text
    path = tmp_path / 'literal.xlsx'
    generator.create_literature_comparison(**data, output_path=path)
    wb = load_workbook(path, data_only=False)
    for sheet, coordinate in zip(wb.worksheets, ['B2', 'B2', 'B2']):
        assert sheet[coordinate].value == text
        assert sheet[coordinate].data_type == 's'
    assert wb.worksheets[0]['B1'].data_type == 's'
    wb.close()


@pytest.mark.parametrize('invalid', ['dimension_count', 'dimension_width', 'gap_width', 'one_paper', 'long_text', 'control', 'sheet_title'])
def test_invalid_data_does_not_produce_workbook(tmp_path, invalid):
    data = sample()
    if invalid == 'dimension_count': data['dimensions'].pop()
    elif invalid == 'dimension_width': data['dimensions'][0].pop()
    elif invalid == 'gap_width': data['gaps'][0]['cells'].pop()
    elif invalid == 'one_paper': data['papers'].pop()
    elif invalid == 'long_text': data['dimensions'][0][0] = 'x' * 32768
    elif invalid == 'control': data['dimensions'][0][0] = 'a\x00b'
    elif invalid == 'sheet_title': data['implication_sheet_title'] = '文獻統整比較表'
    path = tmp_path / 'invalid.xlsx'
    with pytest.raises(ValueError):
        generator.create_literature_comparison(**data, output_path=path)
    assert not path.exists()


def test_existing_manual_workbook_is_preserved(tmp_path):
    path = tmp_path / 'manual.xlsx'
    original = b'previous user file'
    path.write_bytes(original)
    with pytest.raises(FileExistsError):
        generator.create_literature_comparison(**sample(), output_path=path)
    assert path.read_bytes() == original


def test_cli_and_partial_evidence_allow_fewer_gaps(tmp_path):
    data = sample()
    data['gaps'] = []
    source = tmp_path / '比較資料.json'
    source.write_text(json.dumps(data, ensure_ascii=False), encoding='utf-8')
    path = tmp_path / '新成果' / '文獻統整比較表.xlsx'
    result = subprocess.run([sys.executable, str(SCRIPT), '--input', str(source), '--output', str(path)], capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
    wb = load_workbook(path)
    assert wb.worksheets[1].max_row == 1
    wb.close()
